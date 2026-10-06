from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'delivery'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.63)
sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=Inches(.23);sec.footer_distance=Inches(.27)
FONT='Noto Sans CJK SC'
for name,size in [('Normal',10.5),('Title',23),('Subtitle',11),('Heading 1',15),('Heading 2',12),('Caption',8.5)]:
 s=doc.styles[name]; s.font.name=FONT;s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.line_spacing=Pt(15 if name=='Normal' else 11 if name=='Caption' else 28 if name=='Title' else 18 if name=='Heading 1' else 15)
 s.font.italic=False
 if name=='Caption':s.font.bold=False
 for border in list(s.element.xpath('.//w:pBdr')):border.getparent().remove(border)
for name in ['Heading 1','Heading 2']:
 doc.styles[name].font.bold=True
 doc.styles[name].paragraph_format.space_before=Pt(11)
 doc.styles[name].paragraph_format.keep_with_next=True
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
# Compact page numbering, useful for an audit report.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('SAR16 SKY130HS  |  2026-10-06  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

def p(text,style=None,boldlead=None):
 x=doc.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):
  x.add_run(boldlead).bold=True;x.add_run(text[len(boldlead):])
 else:x.add_run(text)
 return x

def h(text):return p(text,'Heading 1')
def h2(text):return p(text,'Heading 2')
def cap(text):
 x=p(text,'Caption');x.paragraph_format.space_after=Pt(7);return x

def table(headers, rows, widths, size=9.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cs=t.add_row().cells
  for i,x in enumerate(row):cs[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:
   rep=OxmlElement('w:tblHeader');trPr.append(rep)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr()
   bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   tcPr.append(bd)
   mar=OxmlElement('w:tcMar')
   for side,value in [('top',65),('bottom',65),('left',95),('right',95)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EEF5' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));tcPr.append(sh)
   for x in c.paragraphs:
    x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=Pt(13)
    if ci>0:x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in x.runs:r.font.size=Pt(size);r.bold=ri==0
 sp=p('');sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3);sp.paragraph_format.space_before=Pt(0)
 return t

def image(name,width):
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER
 x.paragraph_format.space_after=Pt(3);x.paragraph_format.line_spacing=1.0
 x.add_run().add_picture(str(ROOT/'candidate_v2_diodefree/images'/name),width=Inches(width))
 # Native image metadata for accessible editable output.
 for d in x._p.xpath('.//wp:docPr'):d.set('descr',name)
 return x

def page():
 x=doc.add_page_break();x.paragraph_format.line_spacing=Pt(1);x.paragraph_format.space_after=Pt(0)
def mono(text):
 x=p(text);x.paragraph_format.space_after=Pt(4)
 for r in x.runs:r.font.name='Noto Sans Mono';r.font.size=Pt(8)
 return x

# PAGE 1
p('SAR16 SKY130HS\nDRC 与 LVS 最终修复报告','Title')
p('2026 年 10 月 6 日  |  最终交付版本 candidate_v2_diodefree','Subtitle')
p('最终候选已通过本次执行的全部开源检查：真实合并 GDS 的 Magic full DRC 为 0，KLayout BEOL 与 off-grid 为 0，晶体管级提取 LVS 唯一匹配且器件属性错误为 0。该候选替代先前 run_v6 作为最终交付；原始 run_v6 保留不变。',boldlead='最终候选已通过本次执行的全部开源检查：')
p('验证对象为 SKY130HS 1.8 V 数字核心。结果不等同于专有 foundry 签核或完整模拟 ADC 验证；密度检查及其余覆盖边界见第 4 页。')
h('最终验证结果')
table(['检查项','最终结果'],[
 ('Magic 8.3.684 真实 GDS drc(full)','0 个原始违规矩形'),
 ('KLayout 0.30.0 BEOL / off-grid','0 个标记'),
 ('Netgen 1.5.324 提取 LVS','唯一匹配；属性错误 0'),
 ('LVS 网数 / 功能单元数','4,335 / 4,335；4,301 / 4,301'),
 ('Fresh-process 天线检查','net 0 / pin 0；diode 0'),
 ('Router DRC 布局及信号与电源连接','PASS'),
 ('SPEF 已连接网覆盖 电气 CDC 与复位','PASS'),
 ('数字等价及负对照','9 项检查，两次运行全部 PASS')
],[3.8,3.26])
h('时序边界保持不变')
p('时钟约束仍为 clk 10 ns、dec_clk 3 ns，输出负载仍为 50 fF。两组模型下 setup 与 hold 最差裕量均为正，13 项 CDC bounds、Gray span 和复位检查通过。')
table(['模型','Setup WNS  ns','Hold WNS  ns'],[
 ('Nominal','1.062018','0.153111'),
 ('Scalar stress','0.947274','0.050880')
],[2.86,2.10,2.10],10)
p('全程未引入规则豁免、金参考改值或器件属性检查删除。',boldlead='全程未引入规则豁免、金参考改值或器件属性检查删除。')
cap('证据：signoff_summary.json；verification/candidate_v2_diodefree/reports；candidate_v2_diodefree/reports/final_summary.json')

# PAGE 2
page();h('失败根因与实际物理修复')
p('冻结原版检查产生 2,838 个 Magic 原始违规矩形，其中含 nwell 间距与宽度、tap 连通/距离及 LI 间距问题。矩形数不等于独立根因数。提取布局有 4,577 个网，参考有 4,335 个网；多出的 242 个均为未接回 VDD 的 VPB/well 岛。')
p('缺少标准单元 filler 导致 well 不连续；另有 3 个 dfrtp_1 时钟接入点，LEF 的 CLK 矩形覆盖实际 GDS 缺口，使 mcon 接入靠近缺口，产生 6 个 KLayout LI 间距标记。')
h2('执行的修复')
for text in [
 '1  插入 8,900 个官方 HS fill_1/2/4/8 单元，填补原有空位并恢复 well 连续性。',
 '2  将 3 个 CLK 接入 mcon 的局部 y 坐标从 1.665 移至 1.480 µm，各补 0.185 µm metal1 小折线。',
 '3  移除 8 个 antenna diode，并补回曾借用其内部 LI 连通的真实布线分支。',
 '4  用 3 处实际跨层布线解决天线：clk 的 1.44 µm met4 跨接、时钟末端 0.72 µm 分支由 met3 升到 met4，以及 _0827_ 的 met2 桥接，并补齐合法 landing 面积。'
]:p(text)
h2('时钟跨接的真实 GDS 几何')
image('jumper_clk.png',6.9)
cap('图 1  左侧为 candidate v1，右侧为最终无 diode 候选。图像来自实际 GDS 多边形；右侧新增 met4 跨接和两处 M3–M4 接触。坐标单位为 µm。')
p('最终冻结 ODB 在新进程中重新读取后，天线违规为 net 0 / pin 0，diode 实例和 DIODE 端子均为 0。全部 3,230 个 master-pin 的 gate/diffusion 天线模型与工艺天线规则保持原值，未通过放宽模型取得通过结果。')
cap('证据：candidate_v2_diodefree/reports 中的 access_edits.json、diodefree_route_edits.json、invariance_audit.json、antenna_model_integrity.json 与 antenna_fresh_full.rpt')

# PAGE 3
page();h('最终版图与面积')
p('外框仍为 392.98 × 392.98 µm，面积 154,433.2804 µm²；core 为 135,032.832 µm²。RTL、保留实例的连接/位置/方向、178 个物理端口与 4,200 个数据库网名和类型保持不变。')
image('01_routed_full.png',4.55)
cap('图 2  最终候选的实际全局布线视图，含标准单元行、信号布线与电源网络。')
table(['类别','实例数','面积  µm²'],[
 ('逻辑含 buffers','4,301','64,567.3680'),
 ('Tap','1,456','2,327.2704'),
 ('Filler','8,900','68,138.1936'),
 ('Antenna diode','0','0'),
 ('合计','14,657','135,032.8320')
],[3.1,1.5,2.46],10)
p('Filler 填补的是既有 core 空位，不代表增加逻辑功能或扩大芯片外框。最终含 98 种 master，空 cell 为 0。')
cap('证据：candidate_v2_diodefree/reports/final_summary.json、invariance_audit.json；实际图像位于 candidate_v2_diodefree/images')

# PAGE 4
page();h('可复核性与覆盖边界')
h2('参考库选择与数字证明')
p('官方 HS diode_2 的实际扩散区为 0.69 × 0.93 µm，Magic 与 KLayout 均得到周长 3.24 µm，而官方 SPICE 为 7.32 µm；独立单元提取复现同一属性差异。最终通过真实布线消除 diode 需求，未修改参考或豁免属性。参考版本与冻结 ORFS GDS 几何匹配，也避开新版 conb_1 的布局与电阻长度变化。')
p('数字等价覆盖 RTL → mapped → 最终 routed，两次各 9 项检查通过，其中 5 项负对照均产生预期反例。每组比较覆盖 20 个功能输出端口共 146 bit，mapped/routed 各 846 个状态 bit；无未解析功能 blackbox。该证明采用默认参数、复位可达状态与理想布尔/Liberty 语义。')
h2('工具与来源版本')
table(['项目','锁定版本'],[
 ('Magic / Netgen / KLayout','8.3.684 / 1.5.324 / 0.30.0'),
 ('open-pdks','1.0.611  ·  801834fc…'),
 ('几何匹配 HS 参考','968e0c3b…'),
 ('ORFS 平台','ad11cc9f…')
],[2.55,4.51],9.4)
cap('完整 commit、官方来源链接与校验和见 pdk/PDK_AUDIT.md、pdk/source_manifest.json 及原环境 RECONSTRUCTION.md；工具来源与 smoke test 见 tools/tool-status.json。')
h2('交付身份绑定')
p('最终交付 GDS 为 candidate_v2_diodefree/results/candidate.gds，其 SHA256 为：')
mono('bc8f2590c514a8f8bfb819057fcf0f848fc8107b0726860d7581b8bcdadd8c4b')
p('独立验证 GDS 除 BGNLIB / BGNSTR 的创建与修改时间外，所有记录与交付 GDS 逐字节一致。对应审计见 verification/candidate_v2_diodefree/reports/final_gds_binding.json。')
h2('仍需由后续集成验证的范围')
for text in [
 '• Magic drc(full) 是本次公开 FEOL/BEOL 几何检查，并非完整专有 foundry rule deck。标准 DRC 不含 density 和完整 antenna signoff。',
 '• 官方密度检查要求至少 700 × 700 µm 窗口；本 392.98 × 392.98 µm block 更小，须在顶层集成后完成窗口密度验证，简单面积比不能替代。',
 '• 天线 0/0 来自未改原始工艺及单元模型的 fresh-process 计算，不构成专有 foundry 天线认证。',
 '• 本结果不证明完整模拟 ADC 的 SNDR / ENOB，也不是 SMIC180 工艺签核。'
]:p(text)
cap('复核入口：README_DRC_LVS.md、signoff_summary.json、scripts/verify_variant.sh；数字证明详见 equivalence_candidate_v2_diodefree/final_summary.json。')

# Neutral metadata.
doc.core_properties.title='SAR16 SKY130HS DRC 与 LVS 最终修复报告'
doc.core_properties.subject='candidate_v2_diodefree 开源物理验证结果'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.keywords='SAR16 SKY130HS DRC LVS 20261006'
path=OUT/'SAR16_SKY130HS_DRC_LVS_Final_20261006.docx'
doc.save(path)
print(path)
